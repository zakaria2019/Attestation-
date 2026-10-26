# -*- coding: utf-8 -*-
"""
Application Windows - Gestion des demandes d'attestations
Aéroport Ouarzazate / ONDA

Dépendances:
    pip install openpyxl pywin32
Pour créer l'exécutable:
    pip install pyinstaller
    pyinstaller --noconfirm --clean --windowed --name "Attestations Ouarzazate" --hidden-import win32timezone app.py
"""
from __future__ import annotations

import os
import sys
import json
import html
import threading
import traceback
from datetime import datetime
from pathlib import Path
import tkinter as tk
from tkinter import ttk, filedialog, messagebox

try:
    from openpyxl import load_workbook, Workbook
except ImportError:
    load_workbook = None
    Workbook = None

try:
    import win32com.client
except ImportError:
    win32com = None


APP_NAME = "Attestations Ouarzazate"
SHEET_PERSONNEL = "Personnel"
SHEET_JOURNAL = "Journal_Demandes"
DEFAULT_RECIPIENT = "i.fadil@onda.ma"
MAIL_DOMAIN = "@onda.ma"

# Palette institutionnelle inspirée de l'univers visuel aéroportuaire
NAVY = "#113D60"
NAVY_DARK = "#0B2942"
BLUE = "#2A7EAA"
GOLD = "#C28B39"
BG = "#F3F6F9"
WHITE = "#FFFFFF"
TEXT = "#253748"
MUTED = "#667889"
BORDER = "#DCE4EB"
PALE_BLUE = "#EAF2F8"
PALE_GOLD = "#FFF5E3"
GREEN = "#25835B"
RED = "#B84040"


def app_folder() -> Path:
    """Dossier où se trouve le script ou l'exécutable."""
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent


CONFIG_PATH = app_folder() / "config_attestations.json"


def load_config() -> dict:
    try:
        if CONFIG_PATH.exists():
            return json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
    except Exception:
        pass
    return {}


def save_config(data: dict) -> None:
    CONFIG_PATH.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def norm(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def safe_mail_part(value: str) -> str:
    return html.escape(value or "", quote=True)


class WorkbookService:
    """Accès au classeur Personnel / Journal_Demandes."""

    JOURNAL_HEADERS = [
        "Date", "Matricule", "Nom", "Prénom", "Type de demande",
        "À", "CC", "Statut"
    ]

    def __init__(self, workbook_path: str):
        self.path = Path(workbook_path)

    def _load(self):
        if load_workbook is None:
            raise RuntimeError("Le module openpyxl manque. Installez-le avec : pip install openpyxl")
        if not self.path.exists():
            raise FileNotFoundError(f"Classeur introuvable : {self.path}")
        keep_vba = self.path.suffix.lower() == ".xlsm"
        return load_workbook(self.path, data_only=True, keep_vba=keep_vba)

    def get_personnel(self) -> list[dict]:
        wb = self._load()
        try:
            if SHEET_PERSONNEL not in wb.sheetnames:
                raise ValueError(f"La feuille « {SHEET_PERSONNEL} » est absente du classeur.")
            ws = wb[SHEET_PERSONNEL]
            people = []
            # Structure attendue d'après le code VBA fourni :
            # A = matricule, B = nom, C = prénom, D = civilité, I = emploi.
            for row in range(2, ws.max_row + 1):
                matricule = norm(ws.cell(row, 1).value)
                if not matricule:
                    continue
                people.append({
                    "matricule": matricule,
                    "nom": norm(ws.cell(row, 2).value),
                    "prenom": norm(ws.cell(row, 3).value),
                    "civilite": norm(ws.cell(row, 4).value),
                    "emploi": norm(ws.cell(row, 9).value),
                })
            return people
        finally:
            wb.close()

    def get_history(self) -> list[dict]:
        wb = self._load()
        try:
            if SHEET_JOURNAL not in wb.sheetnames:
                return []
            ws = wb[SHEET_JOURNAL]
            items = []
            for row in range(2, ws.max_row + 1):
                values = [ws.cell(row, col).value for col in range(1, 9)]
                if not any(v is not None for v in values):
                    continue
                date_value = values[0]
                if isinstance(date_value, datetime):
                    date_text = date_value.strftime("%d/%m/%Y %H:%M")
                else:
                    date_text = norm(date_value)
                items.append({
                    "date": date_text,
                    "matricule": norm(values[1]),
                    "nom": norm(values[2]),
                    "prenom": norm(values[3]),
                    "type": norm(values[4]),
                    "destinataire": norm(values[5]),
                    "cc": norm(values[6]),
                    "statut": norm(values[7]),
                })
            return list(reversed(items))
        finally:
            wb.close()

    def add_history(self, person: dict, request_type: str, recipient: str, cc: str, status: str) -> None:
        # Pour un fichier .xlsm, conserve les macros avec keep_vba=True.
        keep_vba = self.path.suffix.lower() == ".xlsm"
        wb = load_workbook(self.path, keep_vba=keep_vba)
        try:
            if SHEET_JOURNAL not in wb.sheetnames:
                ws = wb.create_sheet(SHEET_JOURNAL)
                ws.append(self.JOURNAL_HEADERS)
                for cell in ws[1]:
                    cell.font = cell.font.copy(bold=True, color="FFFFFF")
                    cell.fill = cell.fill.copy(fill_type="solid", fgColor="113D60")
                widths = {"A": 20, "B": 16, "C": 22, "D": 22, "E": 42, "F": 30, "G": 30, "H": 16}
                for col, width in widths.items():
                    ws.column_dimensions[col].width = width
            ws = wb[SHEET_JOURNAL]
            if ws.max_row == 1 and not any(ws.cell(1, c).value for c in range(1, 9)):
                for idx, header in enumerate(self.JOURNAL_HEADERS, 1):
                    ws.cell(1, idx).value = header
            ws.append([
                datetime.now(), person["matricule"], person["nom"], person["prenom"],
                request_type, recipient, cc, status
            ])
            last = ws.max_row
            ws.cell(last, 1).number_format = "dd/mm/yyyy hh:mm"
            for cell in ws[last]:
                cell.font = cell.font.copy(name="Aptos", size=10, color="253748")
                cell.alignment = cell.alignment.copy(vertical="center")
                if last % 2 == 0:
                    cell.fill = cell.fill.copy(fill_type="solid", fgColor="F5F9FC")
            wb.save(self.path)
        finally:
            wb.close()


class OutlookService:
    @staticmethod
    def prepare_email(person: dict, request_type: str, recipient: str, cc: str, send_direct: bool = False) -> str:
        if sys.platform != "win32":
            raise RuntimeError("La préparation Outlook nécessite Windows avec Outlook installé.")
        try:
            import win32com.client
        except ImportError:
            raise RuntimeError("Le module pywin32 manque. Installez-le avec : pip install pywin32")

        civilite = person.get("civilite", "").strip().lower()
        civ = "Mme" if civilite in ("madame", "mme", "m", "féminin", "feminin") else "M."
        emploi = person.get("emploi", "").strip()
        if emploi:
            emploi = emploi[:1].upper() + emploi[1:].lower()

        if request_type == "Attestation de salaire":
            libelle = "l’attestation de salaire"
            subject = "Demande de l’attestation de salaire _Aéroport Ouarzazate"
        elif request_type == "Attestation de domiciliation de salaire":
            libelle = "l’attestation de domiciliation de salaire"
            subject = "Demande de l’attestation de domiciliation de salaire _Aéroport Ouarzazate"
        else:
            libelle = "l’attestation de salaire et l’attestation de domiciliation de salaire"
            subject = "Demande de l’attestation de salaire et de domiciliation de salaire _Aéroport Ouarzazate"

        lieu = " à l’aéroport Ouarzazate"
        person_line = (
            f"{civ} {person['nom']} {person['prenom']} "
            f"(Matricule : {person['matricule']})"
        )
        if emploi:
            person_line += f", {emploi}{lieu}"
        body = (
            '<div style="font-family:Century Gothic,Arial,sans-serif;font-size:12pt;color:#253748;">'
            '<p>Bonjour,</p>'
            f'<p>Je vous prie de bien vouloir établir {html.escape(libelle)} de '
            f'{html.escape(person_line)}.</p>'
            '<p>Cordialement,</p></div>'
        )

        outlook = win32com.client.Dispatch("Outlook.Application")
        mail = outlook.CreateItem(0)
        mail.To = recipient
        mail.CC = cc
        mail.Subject = subject
        mail.Display()  # Charge la signature Outlook.
        signature = mail.HTMLBody or ""
        body_pos = signature.lower().find("<body")
        if body_pos >= 0:
            body_end = signature.find(">", body_pos)
            if body_end >= 0:
                mail.HTMLBody = signature[:body_end + 1] + body + signature[body_end + 1:]
            else:
                mail.HTMLBody = body + signature
        else:
            mail.HTMLBody = body + signature
        if send_direct:
            mail.Send()
            return "Envoyé"
        return "Préparé"


class AttestationsApp(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title(APP_NAME)
        self.geometry("1180x760")
        self.minsize(1020, 660)
        self.configure(bg=BG)
        self.option_add("*Font", ("Segoe UI", 10))
        self.config_data = load_config()
        self.workbook_path = self.config_data.get("workbook_path", "")
        self.people: list[dict] = []
        self.people_by_id: dict[str, dict] = {}
        self.current_page = "dashboard"
        self.status_var = tk.StringVar(value="Prêt")
        self.selected_person = tk.StringVar()
        self.name_var = tk.StringVar()
        self.firstname_var = tk.StringVar()
        self.civility_var = tk.StringVar()
        self.job_var = tk.StringVar()
        self.cc_var = tk.StringVar()
        self.to_var = tk.StringVar(value=DEFAULT_RECIPIENT)
        self.salary_var = tk.BooleanVar(value=True)
        self.domic_var = tk.BooleanVar(value=False)
        self.search_history_var = tk.StringVar()
        self._style()
        self._build_shell()
        self._load_data(initial=True)

    def _style(self):
        style = ttk.Style(self)
        try:
            style.theme_use("clam")
        except tk.TclError:
            pass
        style.configure("Treeview", background=WHITE, fieldbackground=WHITE, foreground=TEXT,
                        rowheight=30, borderwidth=0, font=("Segoe UI", 9))
        style.configure("Treeview.Heading", background=PALE_BLUE, foreground=NAVY,
                        font=("Segoe UI", 9, "bold"), relief="flat", padding=(8, 9))
        style.map("Treeview", background=[("selected", "#D6EAF7")], foreground=[("selected", NAVY_DARK)])
        style.configure("TCombobox", padding=7)
        style.configure("TEntry", padding=8)

    def _build_shell(self):
        self.sidebar = tk.Frame(self, bg=NAVY_DARK, width=230)
        self.sidebar.pack(side="left", fill="y")
        self.sidebar.pack_propagate(False)

        brand = tk.Frame(self.sidebar, bg=NAVY_DARK)
        brand.pack(fill="x", padx=20, pady=(25, 26))
        tk.Label(brand, text="✈", bg=NAVY_DARK, fg=GOLD, font=("Segoe UI Symbol", 27)).pack(anchor="w")
        tk.Label(brand, text="AIRPORTS OF", bg=NAVY_DARK, fg="#AFC7D9",
                 font=("Segoe UI", 9, "bold")).pack(anchor="w", pady=(8, 0))
        tk.Label(brand, text="MOROCCO", bg=NAVY_DARK, fg=WHITE,
                 font=("Segoe UI", 19, "bold")).pack(anchor="w")
        tk.Label(brand, text="Aéroport Ouarzazate", bg=NAVY_DARK, fg="#B8C9D6",
                 font=("Segoe UI", 9)).pack(anchor="w", pady=(4, 0))

        tk.Frame(self.sidebar, bg="#34546B", height=1).pack(fill="x", padx=18, pady=(0, 16))
        self.nav_buttons = {}
        for key, label, symbol in [
            ("dashboard", "Tableau de bord", "▦"),
            ("form", "Nouvelle demande", "＋"),
            ("history", "Historique", "☷"),
            ("settings", "Paramètres", "⚙"),
        ]:
            btn = tk.Button(self.sidebar, text=f"  {symbol}    {label}", anchor="w",
                            command=lambda k=key: self.show_page(k),
                            bg=NAVY_DARK, fg="#D9E5ED", activebackground=BLUE,
                            activeforeground=WHITE, relief="flat", bd=0,
                            font=("Segoe UI", 10), padx=18, pady=13, cursor="hand2")
            btn.pack(fill="x", padx=10, pady=2)
            self.nav_buttons[key] = btn

        spacer = tk.Frame(self.sidebar, bg=NAVY_DARK)
        spacer.pack(fill="both", expand=True)
        tk.Label(self.sidebar, text="RESSOURCES HUMAINES", bg=NAVY_DARK, fg="#8FA8B9",
                 font=("Segoe UI", 8, "bold")).pack(anchor="w", padx=22, pady=(8, 4))
        tk.Label(self.sidebar, text="Gestion des attestations\nVersion 1.0",
                 bg=NAVY_DARK, fg="#B8C9D6", justify="left",
                 font=("Segoe UI", 8)).pack(anchor="w", padx=22, pady=(0, 20))

        self.main = tk.Frame(self, bg=BG)
        self.main.pack(side="left", fill="both", expand=True)
        self.topbar = tk.Frame(self.main, bg=WHITE, height=64,
                               highlightbackground=BORDER, highlightthickness=1)
        self.topbar.pack(fill="x")
        self.topbar.pack_propagate(False)
        self.page_title = tk.Label(self.topbar, text="Tableau de bord", bg=WHITE, fg=NAVY,
                                   font=("Segoe UI", 16, "bold"))
        self.page_title.pack(side="left", padx=28)
        self.workbook_label = tk.Label(self.topbar, text="Classeur non sélectionné",
                                       bg=WHITE, fg=MUTED, font=("Segoe UI", 9))
        self.workbook_label.pack(side="right", padx=24)
        self.content = tk.Frame(self.main, bg=BG)
        self.content.pack(fill="both", expand=True, padx=26, pady=22)
        self.footer = tk.Label(self.main, textvariable=self.status_var, anchor="w",
                               bg=WHITE, fg=MUTED, font=("Segoe UI", 9),
                               padx=20, pady=8, highlightbackground=BORDER, highlightthickness=1)
        self.footer.pack(fill="x", side="bottom")

    def _load_data(self, initial=False):
        if not self.workbook_path or not Path(self.workbook_path).exists():
            if initial:
                self.after(300, self._ask_workbook)
            self._update_workbook_label()
            return
        try:
            service = WorkbookService(self.workbook_path)
            self.people = service.get_personnel()
            self.people_by_id = {p["matricule"]: p for p in self.people}
            self._update_workbook_label()
            self.status_var.set(f"{len(self.people)} agents chargés depuis le classeur.")
            if self.current_page == "dashboard":
                self.show_page("dashboard")
        except Exception as exc:
            messagebox.showerror("Erreur de lecture", str(exc))
            self.status_var.set("Erreur de lecture du classeur.")

    def _ask_workbook(self):
        path = filedialog.askopenfilename(
            title="Sélectionner le classeur contenant la feuille Personnel",
            filetypes=[("Classeurs Excel", "*.xlsx *.xlsm"), ("Tous les fichiers", "*.*")]
        )
        if path:
            self.workbook_path = path
            self.config_data["workbook_path"] = path
            save_config(self.config_data)
            self._load_data()
        else:
            self.show_page("settings")

    def _update_workbook_label(self):
        if self.workbook_path:
            self.workbook_label.config(text=Path(self.workbook_path).name)
        else:
            self.workbook_label.config(text="Classeur non sélectionné")

    def _clear_content(self):
        for child in self.content.winfo_children():
            child.destroy()

    def _set_active_nav(self, key):
        for k, btn in self.nav_buttons.items():
            btn.configure(bg=BLUE if k == key else NAVY_DARK,
                          fg=WHITE if k == key else "#D9E5ED")

    def show_page(self, page: str):
        self.current_page = page
        self._set_active_nav(page)
        self._clear_content()
        titles = {
            "dashboard": "Tableau de bord",
            "form": "Nouvelle demande d’attestation",
            "history": "Historique des demandes",
            "settings": "Paramètres de l’application",
        }
        self.page_title.config(text=titles.get(page, APP_NAME))
        if page == "dashboard":
            self._dashboard_page()
        elif page == "form":
            self._form_page()
        elif page == "history":
            self._history_page()
        else:
            self._settings_page()

    def _card(self, parent, title, value, accent=BLUE, subtitle=""):
        card = tk.Frame(parent, bg=WHITE, highlightbackground=BORDER, highlightthickness=1)
        tk.Frame(card, bg=accent, height=4).pack(fill="x")
        inner = tk.Frame(card, bg=WHITE)
        inner.pack(fill="both", expand=True, padx=18, pady=14)
        tk.Label(inner, text=title.upper(), bg=WHITE, fg=MUTED,
                 font=("Segoe UI", 8, "bold")).pack(anchor="w")
        tk.Label(inner, text=str(value), bg=WHITE, fg=NAVY,
                 font=("Segoe UI", 23, "bold")).pack(anchor="w", pady=(7, 2))
        if subtitle:
            tk.Label(inner, text=subtitle, bg=WHITE, fg=MUTED,
                     font=("Segoe UI", 8)).pack(anchor="w")
        return card

    def _dashboard_page(self):
        tk.Label(self.content, text="Bonjour,", bg=BG, fg=MUTED,
                 font=("Segoe UI", 11)).pack(anchor="w")
        tk.Label(self.content, text="Espace de gestion des attestations", bg=BG, fg=NAVY,
                 font=("Segoe UI", 21, "bold")).pack(anchor="w", pady=(2, 6))
        tk.Label(self.content, text="Préparez les demandes RH et consultez leur suivi depuis une interface unique.",
                 bg=BG, fg=MUTED, font=("Segoe UI", 10)).pack(anchor="w", pady=(0, 20))

        history = []
        try:
            if self.workbook_path and Path(self.workbook_path).exists():
                history = WorkbookService(self.workbook_path).get_history()
        except Exception:
            pass
        today = datetime.now().strftime("%d/%m/%Y")
        today_count = sum(1 for item in history if item["date"].startswith(today))
        prepared = sum(1 for item in history if item["statut"].lower() == "préparé")
        sent = sum(1 for item in history if item["statut"].lower() == "envoyé")
        cards = tk.Frame(self.content, bg=BG)
        cards.pack(fill="x")
        card_data = [
            ("Agents référencés", len(self.people), BLUE, "Personnel du classeur"),
            ("Demandes aujourd’hui", today_count, GOLD, today),
            ("Courriels préparés", prepared, GREEN, "À vérifier dans Outlook"),
            ("Demandes envoyées", sent, NAVY, "Envoi direct activé le cas échéant"),
        ]
        for i, (title, val, color, sub) in enumerate(card_data):
            cards.columnconfigure(i, weight=1, uniform="cards")
            self._card(cards, title, val, color, sub).grid(row=0, column=i, sticky="nsew",
                                                           padx=(0 if i == 0 else 8, 0 if i == 3 else 8))

        actions = tk.Frame(self.content, bg=WHITE, highlightbackground=BORDER, highlightthickness=1)
        actions.pack(fill="x", pady=22)
        tk.Label(actions, text="ACCÈS RAPIDE", bg=WHITE, fg=NAVY,
                 font=("Segoe UI", 10, "bold")).pack(anchor="w", padx=20, pady=(18, 12))
        row = tk.Frame(actions, bg=WHITE)
        row.pack(fill="x", padx=20, pady=(0, 20))
        self._action_button(row, "＋  Nouvelle demande", self.show_new_form, BLUE).pack(side="left", padx=(0, 10))
        self._action_button(row, "☷  Consulter l’historique", lambda: self.show_page("history"), NAVY).pack(side="left", padx=10)
        self._action_button(row, "↻  Actualiser les données", self.refresh_data, "#536B7D").pack(side="left", padx=10)

        recent = tk.Frame(self.content, bg=WHITE, highlightbackground=BORDER, highlightthickness=1)
        recent.pack(fill="both", expand=True)
        head = tk.Frame(recent, bg=WHITE)
        head.pack(fill="x", padx=18, pady=14)
        tk.Label(head, text="Demandes récentes", bg=WHITE, fg=NAVY,
                 font=("Segoe UI", 11, "bold")).pack(side="left")
        tk.Button(head, text="Voir tout", command=lambda: self.show_page("history"),
                  bg=WHITE, fg=BLUE, relief="flat", cursor="hand2").pack(side="right")
        tree = self._make_history_tree(recent, compact=True)
        tree.pack(fill="both", expand=True, padx=12, pady=(0, 12))
        for item in history[:5]:
            tree.insert("", "end", values=(item["date"], item["matricule"],
                                           f'{item["nom"]} {item["prenom"]}',
                                           item["type"], item["statut"]))

    def _action_button(self, parent, text, command, color):
        return tk.Button(parent, text=text, command=command, bg=color, fg=WHITE,
                         activebackground=NAVY_DARK, activeforeground=WHITE,
                         relief="flat", bd=0, padx=18, pady=12,
                         font=("Segoe UI", 9, "bold"), cursor="hand2")

    def show_new_form(self):
        self.show_page("form")

    def _form_page(self):
        tk.Label(self.content, text="Créez une demande pour un agent référencé dans le classeur RH.",
                 bg=BG, fg=MUTED, font=("Segoe UI", 10)).pack(anchor="w", pady=(0, 16))
        card = tk.Frame(self.content, bg=WHITE, highlightbackground=BORDER, highlightthickness=1)
        card.pack(fill="both", expand=True)
        tk.Frame(card, bg=BLUE, height=4).pack(fill="x")
        form = tk.Frame(card, bg=WHITE)
        form.pack(fill="both", expand=True, padx=28, pady=22)
        tk.Label(form, text="INFORMATIONS DU PERSONNEL", bg=WHITE, fg=NAVY,
                 font=("Segoe UI", 10, "bold")).grid(row=0, column=0, columnspan=4, sticky="w", pady=(0, 14))
        for col in (1, 3):
            form.columnconfigure(col, weight=1)

        self._field_label(form, "Matricule *", 1, 0)
        self.matricule_combo = ttk.Combobox(form, textvariable=self.selected_person, state="readonly")
        self.matricule_combo.grid(row=2, column=0, columnspan=2, sticky="ew", padx=(0, 16), pady=(0, 12), ipady=4)
        self.matricule_combo["values"] = [p["matricule"] for p in self.people]
        self.matricule_combo.bind("<<ComboboxSelected>>", self._on_person_selected)

        self._field_label(form, "Nom", 1, 2)
        self._readonly_field(form, self.name_var, 2, 2)
        self._field_label(form, "Prénom", 3, 0)
        self._readonly_field(form, self.firstname_var, 4, 0, padx=(0, 16))
        self._field_label(form, "Civilité", 3, 2)
        self._readonly_field(form, self.civility_var, 4, 2)

        self._field_label(form, "Emploi", 5, 0)
        self._readonly_field(form, self.job_var, 6, 0, colspan=3)

        tk.Frame(form, bg=BORDER, height=1).grid(row=7, column=0, columnspan=4, sticky="ew", pady=16)
        tk.Label(form, text="PARAMÈTRES DU COURRIEL", bg=WHITE, fg=NAVY,
                 font=("Segoe UI", 10, "bold")).grid(row=8, column=0, columnspan=4, sticky="w", pady=(0, 12))

        self._field_label(form, "Destinataire (À) *", 9, 0)
        ttk.Entry(form, textvariable=self.to_var).grid(row=10, column=0, columnspan=2, sticky="ew", padx=(0, 16), pady=(0, 12), ipady=4)
        self._field_label(form, "Adresse en copie (CC)", 9, 2)
        ttk.Entry(form, textvariable=self.cc_var).grid(row=10, column=2, columnspan=2, sticky="ew", pady=(0, 12), ipady=4)

        tk.Label(form, text="ATTESTATIONS À DEMANDER", bg=WHITE, fg=NAVY,
                 font=("Segoe UI", 10, "bold")).grid(row=11, column=0, columnspan=4, sticky="w", pady=(8, 10))
        tk.Checkbutton(form, text="Attestation de salaire", variable=self.salary_var,
                       bg=WHITE, activebackground=WHITE, fg=TEXT, selectcolor=PALE_BLUE,
                       font=("Segoe UI", 10)).grid(row=12, column=0, columnspan=2, sticky="w")
        tk.Checkbutton(form, text="Attestation de domiciliation de salaire", variable=self.domic_var,
                       bg=WHITE, activebackground=WHITE, fg=TEXT, selectcolor=PALE_BLUE,
                       font=("Segoe UI", 10)).grid(row=12, column=2, columnspan=2, sticky="w")

        bottom = tk.Frame(form, bg=WHITE)
        bottom.grid(row=13, column=0, columnspan=4, sticky="ew", pady=(22, 0))
        tk.Label(bottom, text="Le courriel s’ouvrira dans Outlook pour vérification avant envoi.",
                 bg=WHITE, fg=MUTED, font=("Segoe UI", 9)).pack(side="left")
        self._action_button(bottom, "Préparer le courriel  →", self.submit_request, BLUE).pack(side="right")

        if self.selected_person.get() in self.people_by_id:
            self._on_person_selected()

    def _field_label(self, parent, text, row, col):
        tk.Label(parent, text=text, bg=WHITE, fg=MUTED,
                 font=("Segoe UI", 9, "bold")).grid(row=row, column=col, sticky="w", pady=(0, 6))

    def _readonly_field(self, parent, var, row, col, colspan=1, padx=(0, 16)):
        entry = tk.Entry(parent, textvariable=var, state="readonly", readonlybackground="#F2F6F9",
                         fg=TEXT, relief="flat", highlightthickness=1,
                         highlightbackground=BORDER, highlightcolor=BLUE)
        entry.grid(row=row, column=col, columnspan=colspan, sticky="ew",
                   padx=padx if col in (0, 2) else (0, 0), pady=(0, 12), ipady=8)

    def _on_person_selected(self, event=None):
        mat = self.selected_person.get().strip()
        person = self.people_by_id.get(mat)
        if not person:
            return
        self.name_var.set(person["nom"])
        self.firstname_var.set(person["prenom"])
        self.civility_var.set(person["civilite"])
        self.job_var.set(person["emploi"])
        # Reproduit la convention du code VBA : première lettre du prénom + . + nom sans espaces.
        if person["nom"] and person["prenom"]:
            self.cc_var.set(f"{person['prenom'][0].lower()}.{person['nom'].lower().replace(' ', '')}{MAIL_DOMAIN}")

    def submit_request(self):
        mat = self.selected_person.get().strip()
        recipient = self.to_var.get().strip()
        cc = self.cc_var.get().strip()
        if not mat or mat not in self.people_by_id:
            messagebox.showwarning("Information manquante", "Veuillez sélectionner un matricule valide.")
            return
        if not recipient:
            messagebox.showwarning("Information manquante", "Le destinataire (À) est obligatoire.")
            return
        if not self.salary_var.get() and not self.domic_var.get():
            messagebox.showwarning("Information manquante", "Sélectionnez au moins un type d’attestation.")
            return
        if self.salary_var.get() and self.domic_var.get():
            request_type = "Salaire + domiciliation"
        elif self.salary_var.get():
            request_type = "Attestation de salaire"
        else:
            request_type = "Attestation de domiciliation de salaire"
        person = self.people_by_id[mat]
        self.status_var.set("Préparation du courriel Outlook…")
        self.update_idletasks()
        try:
            status = OutlookService.prepare_email(person, request_type, recipient, cc, send_direct=False)
            WorkbookService(self.workbook_path).add_history(person, request_type, recipient, cc, status)
            self.status_var.set(f"Demande {status.lower()} et enregistrée dans l’historique.")
            messagebox.showinfo("Demande préparée", "Le courriel a été ouvert dans Outlook pour vérification.\nLa demande a été enregistrée dans l’historique.")
            self.show_page("dashboard")
        except Exception as exc:
            self.status_var.set("La demande n’a pas pu être préparée.")
            messagebox.showerror("Erreur", f"{exc}\n\nVérifiez qu’Outlook est installé et configuré, et que le classeur est accessible.")

    def _make_history_tree(self, parent, compact=False):
        columns = ("date", "matricule", "agent", "type", "statut")
        tree = ttk.Treeview(parent, columns=columns, show="headings", height=6 if compact else 14)
        for col, label, width in [
            ("date", "Date", 145), ("matricule", "Matricule", 95),
            ("agent", "Agent", 190), ("type", "Type de demande", 310), ("statut", "Statut", 100)
        ]:
            tree.heading(col, text=label)
            tree.column(col, width=width, minwidth=70, anchor="w")
        tree.tag_configure("Envoyé", foreground=GREEN)
        tree.tag_configure("Préparé", foreground=BLUE)
        return tree

    def _history_page(self):
        top = tk.Frame(self.content, bg=BG)
        top.pack(fill="x", pady=(0, 14))
        tk.Label(top, text="Consultez les demandes préparées ou envoyées depuis cette application.",
                 bg=BG, fg=MUTED).pack(side="left")
        ttk.Entry(top, textvariable=self.search_history_var, width=28).pack(side="right", ipady=5)
        tk.Button(top, text="Rechercher", command=self._refresh_history_tree, bg=WHITE, fg=NAVY,
                  relief="flat", padx=12, pady=7, cursor="hand2").pack(side="right", padx=8)
        card = tk.Frame(self.content, bg=WHITE, highlightbackground=BORDER, highlightthickness=1)
        card.pack(fill="both", expand=True)
        self.history_tree = self._make_history_tree(card)
        yscroll = ttk.Scrollbar(card, orient="vertical", command=self.history_tree.yview)
        self.history_tree.configure(yscrollcommand=yscroll.set)
        self.history_tree.pack(side="left", fill="both", expand=True, padx=(12, 0), pady=12)
        yscroll.pack(side="right", fill="y", padx=(0, 12), pady=12)
        self._refresh_history_tree()

    def _refresh_history_tree(self):
        if not hasattr(self, "history_tree") or not self.history_tree.winfo_exists():
            return
        for item in self.history_tree.get_children():
            self.history_tree.delete(item)
        try:
            history = WorkbookService(self.workbook_path).get_history() if self.workbook_path else []
        except Exception as exc:
            messagebox.showerror("Historique indisponible", str(exc))
            return
        query = self.search_history_var.get().strip().lower()
        for item in history:
            agent = f'{item["nom"]} {item["prenom"]}'
            searchable = " ".join([item["date"], item["matricule"], agent, item["type"], item["statut"]]).lower()
            if query and query not in searchable:
                continue
            self.history_tree.insert("", "end",
                values=(item["date"], item["matricule"], agent, item["type"], item["statut"]),
                tags=(item["statut"],))

    def _settings_page(self):
        tk.Label(self.content, text="Configuration de la source de données", bg=BG, fg=NAVY,
                 font=("Segoe UI", 18, "bold")).pack(anchor="w", pady=(0, 6))
        tk.Label(self.content, text="Sélectionnez le classeur qui contient la feuille Personnel. Le chemin est mémorisé sur ce PC.",
                 bg=BG, fg=MUTED, wraplength=780, justify="left").pack(anchor="w", pady=(0, 18))
        card = tk.Frame(self.content, bg=WHITE, highlightbackground=BORDER, highlightthickness=1)
        card.pack(fill="x")
        inner = tk.Frame(card, bg=WHITE)
        inner.pack(fill="x", padx=22, pady=22)
        tk.Label(inner, text="CLASSEUR EXCEL RH", bg=WHITE, fg=NAVY,
                 font=("Segoe UI", 10, "bold")).pack(anchor="w", pady=(0, 8))
        self.path_var = tk.StringVar(value=self.workbook_path)
        row = tk.Frame(inner, bg=WHITE)
        row.pack(fill="x")
        ttk.Entry(row, textvariable=self.path_var).pack(side="left", fill="x", expand=True, ipady=6)
        self._action_button(row, "Parcourir…", self._browse_workbook, BLUE).pack(side="left", padx=(10, 0))
        self._action_button(inner, "Enregistrer et charger", self._save_workbook_setting, NAVY).pack(anchor="e", pady=(16, 0))
        note = (
            "Structure attendue dans la feuille Personnel :\n"
            "A = Matricule   |   B = Nom   |   C = Prénom   |   D = Civilité   |   I = Emploi\n\n"
            "La feuille Journal_Demandes est créée automatiquement lors de la première demande si elle n’existe pas.\n"
            "Les demandes sont enregistrées avec le statut « Préparé » lorsque le courriel est ouvert dans Outlook."
        )
        tk.Label(self.content, text=note, bg=BG, fg=MUTED, justify="left",
                 font=("Segoe UI", 9), wraplength=900).pack(anchor="w", pady=20)

    def _browse_workbook(self):
        path = filedialog.askopenfilename(
            title="Choisir le classeur RH",
            filetypes=[("Classeurs Excel", "*.xlsx *.xlsm"), ("Tous les fichiers", "*.*")]
        )
        if path:
            self.path_var.set(path)

    def _save_workbook_setting(self):
        path = self.path_var.get().strip()
        if not path or not Path(path).exists():
            messagebox.showwarning("Classeur invalide", "Choisissez un classeur Excel existant.")
            return
        try:
            service = WorkbookService(path)
            people = service.get_personnel()
            self.workbook_path = path
            self.people = people
            self.people_by_id = {p["matricule"]: p for p in people}
            self.config_data["workbook_path"] = path
            save_config(self.config_data)
            self._update_workbook_label()
            self.status_var.set(f"Classeur chargé : {len(people)} agents.")
            messagebox.showinfo("Configuration enregistrée", f"Classeur chargé avec succès.\nAgents trouvés : {len(people)}")
            self.show_page("dashboard")
        except Exception as exc:
            messagebox.showerror("Erreur de configuration", str(exc))

    def refresh_data(self):
        self._load_data()
        self.status_var.set("Données actualisées.")

if __name__ == "__main__":
    app = AttestationsApp()
    app.mainloop()
